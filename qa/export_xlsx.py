"""Export a feed_verify run (results.json, or partial.jsonl if the run is mid-flight) to a single
.xlsx with one row per banner, its verdict/detail columns, and its banner image embedded in the row
(not an =IMAGE() formula - confirmed unsupported in the user's Excel).

Run:  python -m qa.export_xlsx runs/<stamp>_feedverify [--out runs/<stamp>_feedverify/results.xlsx]
"""
from __future__ import annotations

import argparse
import io
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from PIL import Image as PILImage

THUMB_HEIGHT_PX = 90   # [A] thumbnail height in the sheet; wide enough to read a hero banner, small enough to keep the file light
ROW_HEIGHT_PT = 70     # [A] Excel row height (points) tall enough for THUMB_HEIGHT_PX

COLUMNS = ["image", "result", "reason", "alt_text", "destination_raw", "banner_brands", "banner_deal",
           "listing_title", "title_matches_deal", "missing_brands", "banner_gender", "listing_genders",
           "gender_matches", "total_results", "brands_in_filter", "asset_set", "hotspot_results", "banner_id"]
HEADERS = ["Banner", "Result", "Reason", "Alt text", "Destination", "Banner brands", "Banner deal",
           "Listing title", "Deal matches title?", "Missing brands", "Banner gender", "Listing genders",
           "Gender matches?", "Total results", "Brands in filter", "Asset set", "Hotspot links", "Banner ID"]


def load_results(run_dir: Path) -> list[dict]:
    """Prefer the finished results.json; fall back to partial.jsonl (latest record per banner) mid-run,
    or a run that was cancelled/failed before ever reaching write_outputs(). Neither file exists yet if
    the run died before checking a single banner (e.g. the initial feed fetch itself failed) - treated
    as zero results, not an error, since there's nothing to export."""
    return [_shown(r) for r in _load_raw(run_dir)]


def _shown(r: dict) -> dict:
    """The result as a person sees it: a spent-retries temporary error is UNAVAILABLE, not INCONCLUSIVE."""
    from qa.feed_verify import shown_result
    label = shown_result(r)
    return r if label == r["result"] else {**r, "result": label}


def _load_raw(run_dir: Path) -> list[dict]:
    results_path = run_dir / "results.json"
    if results_path.exists():
        from qa.feed_verify import load_final_results       # results.json plus any later per-banner retry
        return list(load_final_results(run_dir).values())
    partial_path = run_dir / "partial.jsonl"
    if not partial_path.exists():
        return []
    latest: dict[str, dict] = {}
    for line in partial_path.read_text(encoding="utf-8").splitlines():
        line = line.strip()
        if not line:
            continue
        try:
            rec = json.loads(line)
        except json.JSONDecodeError:
            continue
        latest[rec.get("banner_id")] = rec
    order = {"PASS": 0, "FAIL": 1, "INCONCLUSIVE": 2, "UNAVAILABLE": 3, "SKIPPED": 4}
    return sorted(latest.values(), key=lambda r: (order.get(r["result"], 9), r.get("banner_id", "")))


def _cell(r: dict, col: str):
    c = r.get("banner_check") or {}
    if col == "reason":
        from qa.reference_check import note
        base = r.get("reason") or _fail_reason(r, c)
        ref = note(r)
        return f"{base}; {ref}" if base and ref and ref not in base else (base or ref)
    if col == "hotspot_results":
        from qa.feed_verify import _hotspot_summary
        return " | ".join(_hotspot_summary(h) for h in r.get("hotspot_checks", []))
    if col in ("banner_brands", "missing_brands", "listing_genders"):
        return " | ".join(c.get(col, []) or [])
    if col in c:
        return c.get(col)
    return r.get(col, "")


def _fail_reason(r: dict, c: dict) -> str:
    if r["result"] != "FAIL":
        return ""
    parts = []
    if c.get("missing_brands"):
        parts.append(f"missing brands {c['missing_brands']}")
    if c.get("title_matches_deal") is False:
        parts.append(f"title {c.get('listing_title')!r} != deal {c.get('banner_deal')!r}")
    if c.get("gender_matches") is False:
        parts.append(f"targets {c.get('banner_gender')!r} but listing genders are {c.get('listing_genders')}")
    return "; ".join(parts)


def _thumb_png_bytes(image_file: str | None) -> bytes | None:
    if not image_file or not Path(image_file).exists():
        return None
    with PILImage.open(image_file) as im:
        im = im.convert("RGB")
        ratio = THUMB_HEIGHT_PX / im.height
        im = im.resize((max(1, int(im.width * ratio)), THUMB_HEIGHT_PX))
        buf = io.BytesIO()
        im.save(buf, format="PNG")
        return buf.getvalue()


def build_workbook(results: list[dict]) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "banners"
    ws.append(HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    fills = {"PASS": "C6EFCE", "FAIL": "FFC7CE", "INCONCLUSIVE": "FFEB9C", "UNAVAILABLE": "DDEBF7", "SKIPPED": "D9D9D9"}
    from openpyxl.styles import PatternFill

    row_i = 2
    for r in results:
        for col_i, col in enumerate(COLUMNS, start=1):
            if col == "image":
                continue
            val = _cell(r, col)
            cell = ws.cell(row=row_i, column=col_i, value=val if not isinstance(val, bool) else str(val))
            cell.alignment = Alignment(vertical="center", wrap_text=col in ("reason", "destination_raw"))
            if col == "result" and val in fills:
                cell.fill = PatternFill("solid", fgColor=fills[val])
        png = _thumb_png_bytes(r.get("image_file"))
        if png:
            # already resized to THUMB_HEIGHT_PX tall (correct aspect) by _thumb_png_bytes; use as-is
            ws.add_image(XLImage(io.BytesIO(png)), f"A{row_i}")
        ws.row_dimensions[row_i].height = ROW_HEIGHT_PT
        row_i += 1

    ws.column_dimensions["A"].width = 18
    widths = {"result": 13, "reason": 45, "alt_text": 22, "destination_raw": 40, "banner_brands": 28,
              "banner_deal": 20, "listing_title": 22, "title_matches_deal": 10, "missing_brands": 20,
              "banner_gender": 12, "listing_genders": 18, "gender_matches": 10, "total_results": 10,
              "brands_in_filter": 10, "asset_set": 10, "banner_id": 26}
    for col_i, col in enumerate(COLUMNS[1:], start=2):
        ws.column_dimensions[get_column_letter(col_i)].width = widths.get(col, 16)
    ws.freeze_panes = "A2"
    return wb


def main() -> None:
    ap = argparse.ArgumentParser(description="Export a feed_verify run to .xlsx with embedded banner images")
    ap.add_argument("run_dir", type=Path)
    ap.add_argument("--out", type=Path, default=None)
    args = ap.parse_args()
    results = load_results(args.run_dir)
    wb = build_workbook(results)
    out = args.out or (args.run_dir / "results.xlsx")
    wb.save(out)
    print(f"{len(results)} banners -> {out}")


if __name__ == "__main__":
    main()
