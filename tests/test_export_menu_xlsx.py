"""qa/export_menu_xlsx.py: the menu workbook (pictures, alt text, title, link in their own columns)."""
import re
import zipfile

import openpyxl
from PIL import Image

from qa.export_menu_xlsx import build_workbook


def png(path, size=(200, 120)):
    path.parent.mkdir(parents=True, exist_ok=True)
    Image.new("RGB", size, (200, 30, 30)).save(path, "PNG")


def item(i, parent, level, title, files=(), link="", **kw):
    return {"id": i, "parent": parent, "level": level, "title": title, "link": link, "opens": kw.pop("opens", ""), "description": kw.pop("description", ""),
            "alt": kw.pop("alt", ""), "images": [f"https://cdn/{f}" for f in files], "image_files": list(files), "active": kw.pop("active", True), "audience": "", **kw}


def run_dir(tmp_path):
    for name in ("men.png", "shoes.png", "ad.png"):
        png(tmp_path / "images" / name)
    (tmp_path / "images" / "logo.svg").write_text("<svg/>", encoding="utf-8")
    return tmp_path


def test_a_navigation_sheet_has_the_picture_title_link_and_path_in_separate_columns(tmp_path):
    d = run_dir(tmp_path)
    menu = {"kind": "top-nav", "l1": "premium", "state": "KARNATAKA", "fetched_at": "2026-10-05T11:40:00+00:00", "items": [
        item(0, None, 0, "Men", ["men.png"], opens="app screen: /sections/men-nav-page"),
        item(1, 0, 1, "Footwear", ["shoes.png", "logo.svg"], link="https://www.ajio.com/shop/footwear", opens="web page", active=False),
        item(2, 0, 1, "No picture", ["missing.png"])]}
    out = build_workbook([{"menu": menu, "dir": d}], tmp_path / "x.xlsx")
    wb = openpyxl.load_workbook(out)
    assert wb.sheetnames == ["Read me", "Top menu"]
    ws = wb["Top menu"]
    assert [c.value for c in ws[1]][:7] == ["Image", "Second image", "Alt text", "Title", "Link", "What it opens", "Path"]
    rows = [[c.value for c in r] for r in ws.iter_rows(min_row=2)]
    assert [r[3] for r in rows] == ["Men", "Footwear", "No picture"]
    assert rows[1][4] == "https://www.ajio.com/shop/footwear" and ws["E3"].hyperlink.target == "https://www.ajio.com/shop/footwear"
    assert rows[0][4] == "app screen: /sections/men-nav-page" and rows[1][6] == "Men > Footwear" and rows[1][7] == "no"
    assert len(ws._images) == 2                          # Men, and Footwear's first picture; the SVG and the missing file are skipped, not fatal
    z = zipfile.ZipFile(out)
    drawings = "".join(z.read(n).decode() for n in z.namelist() if re.match(r"xl/drawings/drawing\d+\.xml", n))
    assert drawings.count("hlinkClick") == 1             # only Footwear's picture has a link to open


def test_ads_list_one_row_per_ad_with_its_slot_and_trending_has_title_and_description(tmp_path):
    d = run_dir(tmp_path)
    ads = {"kind": "ads", "l1": "nontransacted", "state": "ASSAM", "items": [
        item(0, None, 0, "Home screen", opens="1 ad"), item(1, 0, 1, "", ["ad.png"], link="https://www.indiaistore.com/x", opens="web page", description="rank 3"),
        item(2, None, 0, "My account banner", opens="0 ads")]}
    trends = {"kind": "trending", "l1": "premium", "items": [item(0, None, 0, "#NavratriReady", ["men.png"], description="Twirl-ready", opens="in-app search: :relevance:trend:#N")]}
    wb = openpyxl.load_workbook(build_workbook([{"menu": ads, "dir": d}, {"menu": trends, "dir": d}], tmp_path / "y.xlsx"))
    assert wb.sheetnames == ["Read me", "Sponsored ads", "Trending"]
    a = [[c.value for c in r] for r in wb["Sponsored ads"].iter_rows(min_row=2)]
    assert len(a) == 1 and a[0][4] == "https://www.indiaistore.com/x" and a[0][5] == "Home screen" and a[0][6] == "3"
    t = [[c.value for c in r] for r in wb["Trending"].iter_rows(min_row=2)]
    assert t[0][2] == "#NavratriReady" and t[0][3] == "Twirl-ready" and t[0][4].startswith("in-app search")
    readme = " ".join(str(c.value) for r in wb["Read me"].iter_rows() for c in r)
    assert "Home screen (1 ad)" in readme and "My account banner (0 ads)" in readme and "state ASSAM" in readme
