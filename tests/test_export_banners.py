import csv
import json
from pathlib import Path

from qa import export_banners as exp

SAMPLE = Path(__file__).resolve().parent.parent / "analysis" / "traffic_capture" / "home_theme_response_sample.json"


class FakeResp:
    def __init__(self, status, body, ctype):
        self.status_code = status
        self.content = body
        self.headers = {"content-type": ctype}


def fake_getter(status=200, body=b"fake-bytes", ctype="image/webp"):
    return lambda url, headers=None, timeout=None: FakeResp(status, body, ctype)


def test_fetch_image_requests_webp_first():
    seen = {}

    def getter(url, headers=None, timeout=None):
        seen["headers"] = headers
        class Resp:
            status_code = 200
            content = b"x"
            headers = {"content-type": "image/webp"}
        return Resp()

    exp.fetch_image("http://x", getter)
    assert seen["headers"]["Accept"] == exp.IMAGE_ACCEPT


def test_fetch_image_reports_non_ok_and_non_image():
    assert exp.fetch_image("http://x", fake_getter(status=404))[0] == "HTTP 404"
    assert exp.fetch_image("http://x", fake_getter(ctype="text/html"))[0].startswith("not an image")


def test_download_images_writes_files_and_dedupes_by_content(tmp_path):
    result = exp.download_images(["http://a", "http://b"], tmp_path / "images", fake_getter())
    assert result["http://a"][0] == "OK" and result["http://b"][0] == "OK"
    # identical fake content -> identical hash -> same file, written once
    assert result["http://a"][1] == result["http://b"][1]
    assert list((tmp_path / "images").iterdir())


def test_build_rows_matches_column_order_and_has_no_missing_keys():
    theme = {"sections": json.loads(SAMPLE.read_text(encoding="utf-8"))["one_example_section_per_type"]}
    rows = exp.build_rows(theme, images={}, slug="home", fetched_at="2026-09-21T00:00:00+00:00")
    assert rows and all(set(r) == set(exp.COLUMNS) for r in rows)
    floating = next(r for r in rows if r["section_type"] == "floating-widget")
    assert floating["destination_type"] == "PLP" and floating["image_status"] == "not_downloaded"


def test_write_csv_roundtrips_with_bom_for_excel(tmp_path):
    rows = [{c: f"v{c}" for c in exp.COLUMNS}]
    out = tmp_path / "banners.csv"
    exp.write_csv(rows, out)
    assert out.read_bytes().startswith(b"\xef\xbb\xbf")
    with open(out, encoding="utf-8-sig") as f:
        assert list(csv.DictReader(f)) == rows


def test_write_xlsx_embeds_thumbnails_and_skips_missing_or_broken_files(tmp_path):
    from openpyxl import load_workbook
    from PIL import Image

    (tmp_path / "images").mkdir()
    Image.new("RGB", (1024, 110), "red").save(tmp_path / "images" / "good.webp", format="WEBP")
    (tmp_path / "images" / "broken.webp").write_bytes(b"not an image")

    def row(file):
        return {**{c: "" for c in exp.COLUMNS}, "banner_id": "b", "image_file": file}

    rows = [row("images/good.webp"), row("images/broken.webp"), row("images/missing.webp"), row("")]
    embedded = exp.write_xlsx(rows, tmp_path / "banners.xlsx", tmp_path)
    assert embedded == 1

    ws = load_workbook(tmp_path / "banners.xlsx")["banners"]
    assert [c.value for c in ws[1]][:3] == ["image", "position", "banner_id"]
    assert ws.max_row == 5 and len(ws._images) == 1
    assert ws.row_dimensions[2].height > 18


def test_csv_has_no_excel_365_only_formulas():
    assert "image_preview" not in exp.COLUMNS


def test_build_rows_with_hero_scope_only_keeps_hero_candidates():
    theme = {"sections": json.loads(SAMPLE.read_text(encoding="utf-8"))["one_example_section_per_type"]}
    all_rows = exp.build_rows(theme, images={}, slug="home", fetched_at="2026-09-21T00:00:00+00:00", scope="all")
    hero_rows = exp.build_rows(theme, images={}, slug="home", fetched_at="2026-09-21T00:00:00+00:00", scope="hero")
    assert len(hero_rows) < len(all_rows)
    assert all(r["section_type"] == "hybrid-dynamic-banner" and r["block_index"] != "" for r in hero_rows)


def test_export_end_to_end_without_network(tmp_path):
    run_dir = tmp_path / "run"
    run_dir.mkdir()
    theme = {"sections": json.loads(SAMPLE.read_text(encoding="utf-8"))["one_example_section_per_type"]}
    (run_dir / "home.response.json").write_text(json.dumps(theme), encoding="utf-8")
    rows = exp.export(run_dir, tmp_path / "out", getter=fake_getter())
    assert (tmp_path / "out" / "banners.csv").exists()
    assert (tmp_path / "out" / "banners.xlsx").exists()
    assert any(r["image_status"] == "OK" for r in rows)
