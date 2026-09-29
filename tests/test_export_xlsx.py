"""Excel export: one row per banner with its verdict columns and its image embedded (qa/export_xlsx.py)."""
import json

from openpyxl import load_workbook
from PIL import Image as PILImage

from qa import export_xlsx as xl


def real_image_file(tmp_path, name="a.png", w=200, h=100):
    path = tmp_path / name
    PILImage.new("RGB", (w, h), "blue").save(path)
    return str(path)


def result(bid="a", result="PASS", image_file=None, **extra):
    return {"banner_id": bid, "result": result, "alt_text": bid, "destination_raw": f"https://ajio.com/s/{bid}",
           "image_file": image_file, **extra}


def test_load_results_prefers_the_finished_file_over_partial(tmp_path):
    (tmp_path / "results.json").write_text(json.dumps([result("a"), result("b")]), encoding="utf-8")
    (tmp_path / "partial.jsonl").write_text(json.dumps(result("c")) + "\n", encoding="utf-8")
    rows = xl.load_results(tmp_path)
    assert [r["banner_id"] for r in rows] == ["a", "b"]


def test_load_results_falls_back_to_partial_when_the_run_never_finished(tmp_path):
    (tmp_path / "partial.jsonl").write_text(
        json.dumps(result("a", attempts=1)) + "\n" + json.dumps(result("a", attempts=2)) + "\n" +
        json.dumps(result("b")) + "\n", encoding="utf-8")
    rows = xl.load_results(tmp_path)
    by = {r["banner_id"]: r for r in rows}
    assert set(by) == {"a", "b"}
    assert by["a"]["attempts"] == 2   # the later of two partial.jsonl lines for the same banner wins


def test_load_results_is_empty_not_an_error_when_the_run_died_before_writing_anything(tmp_path):
    assert xl.load_results(tmp_path) == []


def test_build_workbook_has_one_row_per_banner_plus_a_header(tmp_path):
    wb = xl.build_workbook([result("a", "PASS"), result("b", "FAIL")])
    ws = wb.active
    assert ws.max_row == 3
    assert ws.cell(row=1, column=2).value == "Result"
    assert ws.cell(row=2, column=2).value == "PASS"
    assert ws.cell(row=3, column=2).value == "FAIL"


def test_build_workbook_embeds_the_banner_image_when_one_exists_on_disk(tmp_path):
    img = real_image_file(tmp_path)
    wb = xl.build_workbook([result("a", image_file=img)])
    ws = wb.active
    assert len(ws._images) == 1


def test_build_workbook_skips_the_image_silently_when_the_file_is_missing(tmp_path):
    wb = xl.build_workbook([result("a", image_file=str(tmp_path / "does-not-exist.png"))])
    ws = wb.active
    assert len(ws._images) == 0
    assert ws.max_row == 2   # the row itself is still written, just without a thumbnail


def test_a_saved_workbook_round_trips(tmp_path):
    img = real_image_file(tmp_path)
    wb = xl.build_workbook([result("a", "FAIL", image_file=img, reason="missing brands ['Nike']")])
    out = tmp_path / "results.xlsx"
    wb.save(out)
    reloaded = load_workbook(out)
    assert reloaded.active.cell(row=2, column=3).value == "missing brands ['Nike']"


def test_a_spent_retries_temporary_error_is_exported_as_unavailable_not_inconclusive(tmp_path):
    rows = [result("a", "INCONCLUSIVE", reason="gave up: listing_fetch_failed: HTTP 400", attempts=5),
            result("b", "INCONCLUSIVE", reason="empty_bounding_box_after_scaling"), result("c", "PASS")]
    (tmp_path / "results.json").write_text(json.dumps(rows), encoding="utf-8")
    got = {r["banner_id"]: r["result"] for r in xl.load_results(tmp_path)}
    assert got == {"a": "UNAVAILABLE", "b": "INCONCLUSIVE", "c": "PASS"}       # b isn't a server/network problem
